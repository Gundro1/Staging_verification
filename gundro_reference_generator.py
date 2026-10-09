"""
================================================================================
GUNDRO QA STAGING REFERENCE GENERATOR (v5)
================================================================================
Purpose: Reads a BisViews seed dataset (.xlsx or .sql), extracts unique brands,
         calculates total branch counts, and splits them equally among QA verifiers
         into an SOP-compliant Excel verification workbook.

Verifiers use the generated workbook to audit https://staging.bisviews.com.

Dependencies:
    pip install openpyxl pandas

Usage:
    # Interactive mode (prompts for input file & verifier names):
    python gundro_reference_generator.py

    # CLI mode (Excel input):
    python gundro_reference_generator.py --file "path/to/dataset.xlsx" --verifiers Grace Abdul Siddik

    # CLI mode (SQL input):
    python gundro_reference_generator.py --file "path/to/dataset.sql" --verifiers Grace Abdul Siddik
================================================================================
"""

import os
import sys
import re
import math
import argparse
from collections import OrderedDict
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout.reconfigure(encoding='utf-8')

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "QA_Reference_Datasets")


def extract_brands_from_excel(excel_path):
    """
    Parses a BisViews Excel seed file (.xlsx) and extracts:
    - Brand name
    - Category & Subcategory
    - Total Branches
    - Country & Country Code
    - Website
    """
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet_name = 'Franchise' if 'Franchise' in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    headers = [str(cell.value or '').strip() for cell in ws[1]]
    
    def get_col_idx(name_candidates):
        for candidate in name_candidates:
            if candidate in headers:
                return headers.index(candidate)
        return None

    fname_idx = get_col_idx(['franchise_name', 'business_name', 'brand_name'])
    cat_idx = get_col_idx(['category'])
    subcat_idx = get_col_idx(['subcategory'])
    web_idx = get_col_idx(['website', 'url'])
    cname_idx = get_col_idx(['country_name', 'country'])
    ccode_idx = get_col_idx(['country_code'])

    if fname_idx is None:
        raise ValueError("Could not find franchise_name or business_name column in Excel file.")

    brand_data = OrderedDict()
    country_name = "Algeria"
    country_code = "DZ"
    total_rows = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not any(row):
            continue
        
        name = str(row[fname_idx] or '').strip()
        if not name:
            continue

        total_rows += 1
        cat = str(row[cat_idx] or 'Uncategorized').strip() if cat_idx is not None and row[cat_idx] else 'Uncategorized'
        subcat = str(row[subcat_idx] or 'Uncategorized').strip() if subcat_idx is not None and row[subcat_idx] else 'Uncategorized'
        web = str(row[web_idx] or '').strip() if web_idx is not None and row[web_idx] else ''
        
        if cname_idx is not None and row[cname_idx]:
            country_name = str(row[cname_idx]).strip()
        if ccode_idx is not None and row[ccode_idx]:
            country_code = str(row[ccode_idx]).strip()

        if name not in brand_data:
            brand_data[name] = {
                'category': cat,
                'subcategory': subcat,
                'website': web,
                'branches': 0
            }
        brand_data[name]['branches'] += 1

    brands = []
    for name, data in brand_data.items():
        brands.append({
            'Brand Name': name,
            'Category': data['category'],
            'Subcategory': data['subcategory'],
            'Total Branches': data['branches'],
            'Website': data['website'],
            'Country': country_name,
            'Country Code': country_code,
        })

    brands.sort(key=lambda b: (b['Category'], b['Subcategory'], b['Brand Name']))
    for i, b in enumerate(brands):
        b['Test ID'] = f"TC{i+1:03d}"

    total_brands = len(brands)
    total_branches = sum(b['Total Branches'] for b in brands)

    return brands, country_name, country_code, total_rows, total_brands, total_branches


def extract_brands_from_sql(sql_path):
    """
    Parses a BisViews SQL seed file (.sql) with robust CTE matching.
    """
    with open(sql_path, 'r', encoding='utf-8', errors='replace') as f:
        content = f.read()

    with_idx = content.find("WITH")
    insert_idx = content.find("INSERT INTO")
    cte_str = content[with_idx:insert_idx] if with_idx != -1 and insert_idx != -1 else content

    cat_cte = {}
    for m in re.finditer(
        r"_c(\d+)\s+AS\s*\(\s*SELECT\s+id\s+FROM\s+business_categories\s+WHERE\s+LOWER\(name\)\s*=\s*LOWER\('([^']+)'\)",
        cte_str, re.IGNORECASE
    ):
        cat_cte[f"_c{m.group(1)}"] = m.group(2)

    subcat_cte = {}
    for m in re.finditer(
        r"_s(\d+)\s+AS\s*\(\s*SELECT\s+COALESCE\s*\(\s*\(\s*SELECT\s+id\s+FROM\s+business_subcategories\s+WHERE\s+LOWER\(name\)\s*=\s*LOWER\('([^']+)'\)\s+AND\s+category_id\s*=\s*\(\s*SELECT\s+id\s+FROM\s+business_categories\s+WHERE\s+LOWER\(name\)\s*=\s*LOWER\('([^']+)'\)",
        cte_str, re.IGNORECASE
    ):
        subcat_cte[f"_s{m.group(1)}"] = {
            'subcategory': m.group(2),
            'category': m.group(3)
        }

    brand_data = OrderedDict()
    row_pattern = re.compile(
        r"\(\s*'((?:[^']|'')*)',\s*"           # name (group 1)
        r"'(?:[^']|'')*',\s*"                   # slug
        r"\(SELECT id FROM (_c\d+)\),\s*"       # category ref (group 2)
        r"\(SELECT id FROM (_s\d+)\),\s*"       # subcategory ref (group 3)
        r"'seeded',\s*"                         # source
        r"(?:TRUE|FALSE),\s*"                   # is_claimed
        r"(?:'((?:[^']|'')*)'|NULL)"            # website (group 4) or NULL
    )

    country_name = 'Algeria'
    country_code = 'DZ'
    total_rows = 0

    for m in row_pattern.finditer(content):
        total_rows += 1
        name = m.group(1).replace("''", "'")
        cat_ref = m.group(2)
        subcat_ref = m.group(3)
        website = (m.group(4) or '').replace("''", "'")

        category = subcat_cte.get(subcat_ref, {}).get('category', cat_cte.get(cat_ref, 'Uncategorized'))
        subcategory = subcat_cte.get(subcat_ref, {}).get('subcategory', 'Uncategorized')

        if name not in brand_data:
            brand_data[name] = {
                'category': category,
                'subcategory': subcategory,
                'website': website,
                'branches': 0
            }
        brand_data[name]['branches'] += 1

    brands = []
    for name, data in brand_data.items():
        brands.append({
            'Brand Name': name,
            'Category': data['category'],
            'Subcategory': data['subcategory'],
            'Total Branches': data['branches'],
            'Website': data['website'],
            'Country': country_name,
            'Country Code': country_code,
        })

    brands.sort(key=lambda b: (b['Category'], b['Subcategory'], b['Brand Name']))
    for i, b in enumerate(brands):
        b['Test ID'] = f"TC{i+1:03d}"

    total_brands = len(brands)
    total_branches = sum(b['Total Branches'] for b in brands)

    return brands, country_name, country_code, total_rows, total_brands, total_branches


def assign_brands(brands, verifier_names):
    """
    Distributes brands to verifiers using round-robin logic.
    Each brand is assigned to exactly ONE verifier (no sharing).
    """
    num_v = len(verifier_names)
    loads = {name: 0 for name in verifier_names}

    for i, brand in enumerate(brands):
        v_name = verifier_names[i % num_v]
        brand['Assigned Verifier'] = v_name
        loads[v_name] += 1

    return brands, loads


def generate_workbook(brands, verifier_names, loads, country_name, country_code,
                      total_rows, total_brands, total_branches, source_file):
    """Generates an SOP-compliant Excel verification workbook."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"QA_Verification_{country_name}_{country_code}_{timestamp}.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    navy = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    navy_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    blue_bg = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    blue_font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    bold_font = Font(name="Calibri", size=11, bold=True, color="333333")
    reg_font = Font(name="Calibri", size=11, color="000000")
    border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

    num_v = len(verifier_names)
    all_cats = sorted(set(b['Category'] for b in brands))
    all_subcats = sorted(set(b['Subcategory'] for b in brands))

    # MASTER OVERVIEW SHEET
    ws = wb.create_sheet("Master Overview")
    r = 1
    ws.cell(row=r, column=1, value=f"BISVIEWS QA STAGING VERIFICATION — {country_name.upper()} ({country_code})").font = title_font
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 2

    kpis = [
        ("Country:", country_name),
        ("Country Code:", country_code),
        ("Source File:", os.path.basename(source_file)),
        ("Total Branch Rows:", total_rows),
        ("Total Unique Brands:", total_brands),
        ("Total Branch Locations:", total_branches),
        ("Number of Categories:", len(all_cats)),
        ("Number of Subcategories:", len(all_subcats)),
        ("Number of Verifiers:", num_v),
        ("Brands Per Verifier (approx):", math.ceil(total_brands / num_v)),
        ("Generated On:", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Staging URL:", "https://staging.bisviews.com"),
    ]
    for label, value in kpis:
        ws.cell(row=r, column=1, value=label).font = blue_font
        ws.cell(row=r, column=1).fill = blue_bg
        ws.cell(row=r, column=2, value=value).font = bold_font
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="FULL CATEGORY & SUBCATEGORY BREAKDOWN").font = title_font
    r += 1
    for c, h in enumerate(["Category", "Subcategory", "Brand Count", "Total Branches"], 1):
        ws.cell(row=r, column=c, value=h).font = navy_font
        ws.cell(row=r, column=c).fill = navy
        ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")
    r += 1

    cat_sub_data = {}
    for b in brands:
        key = (b['Category'], b['Subcategory'])
        if key not in cat_sub_data:
            cat_sub_data[key] = {'count': 0, 'branches': 0}
        cat_sub_data[key]['count'] += 1
        cat_sub_data[key]['branches'] += b['Total Branches']

    for (cat, subcat), data in sorted(cat_sub_data.items()):
        ws.cell(row=r, column=1, value=cat).font = reg_font
        ws.cell(row=r, column=1).border = border
        ws.cell(row=r, column=2, value=subcat).font = reg_font
        ws.cell(row=r, column=2).border = border
        ws.cell(row=r, column=3, value=data['count']).font = reg_font
        ws.cell(row=r, column=3).border = border
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=4, value=data['branches']).font = reg_font
        ws.cell(row=r, column=4).border = border
        ws.cell(row=r, column=4).alignment = Alignment(horizontal="center")
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = bold_font
    ws.cell(row=r, column=3, value=total_brands).font = bold_font
    ws.cell(row=r, column=3).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=4, value=total_branches).font = bold_font
    ws.cell(row=r, column=4).alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in col if len(str(cell.value or '')) < 60), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 14)

    # INDIVIDUAL VERIFIER SHEETS
    for v_name in verifier_names:
        v_brands = [b for b in brands if b['Assigned Verifier'] == v_name]
        v_count = len(v_brands)
        v_branches = sum(b['Total Branches'] for b in v_brands)
        v_cats = sorted(set(b['Category'] for b in v_brands))
        v_subcats = sorted(set(b['Subcategory'] for b in v_brands))

        ws_v = wb.create_sheet(v_name[:31])
        r = 1
        ws_v.cell(row=r, column=1, value=f"QA STAGING VERIFICATION SHEET — {v_name.upper()}").font = title_font
        ws_v.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        r += 2

        v_kpis = [
            ("Verifier Name:", v_name),
            ("Country:", f"{country_name} ({country_code})"),
            ("Total Brands Assigned:", v_count),
            ("Total Branches to Verify:", v_branches),
            ("Categories Assigned:", len(v_cats)),
            ("Subcategories Assigned:", len(v_subcats)),
            ("Staging URL:", "https://staging.bisviews.com"),
        ]
        for label, value in v_kpis:
            ws_v.cell(row=r, column=1, value=label).font = blue_font
            ws_v.cell(row=r, column=1).fill = blue_bg
            ws_v.cell(row=r, column=2, value=value).font = bold_font
            r += 1
        r += 1

        ws_v.cell(row=r, column=1, value="YOUR CATEGORIES & SUBCATEGORIES").font = title_font
        r += 1
        for c, h in enumerate(["Category", "Subcategory", "Brand Count"], 1):
            ws_v.cell(row=r, column=c, value=h).font = navy_font
            ws_v.cell(row=r, column=c).fill = navy
            ws_v.cell(row=r, column=c).alignment = Alignment(horizontal="center")
        r += 1

        v_cat_sub = {}
        for b in v_brands:
            key = (b['Category'], b['Subcategory'])
            v_cat_sub[key] = v_cat_sub.get(key, 0) + 1

        for (cat, subcat), count in sorted(v_cat_sub.items()):
            ws_v.cell(row=r, column=1, value=cat).font = reg_font
            ws_v.cell(row=r, column=1).border = border
            ws_v.cell(row=r, column=2, value=subcat).font = reg_font
            ws_v.cell(row=r, column=2).border = border
            ws_v.cell(row=r, column=3, value=count).font = reg_font
            ws_v.cell(row=r, column=3).border = border
            ws_v.cell(row=r, column=3).alignment = Alignment(horizontal="center")
            r += 1
        r += 1

        ws_v.cell(row=r, column=1, value="Filter by Category:").font = blue_font
        ws_v.cell(row=r, column=1).fill = blue_bg
        if v_cats:
            v_cat_str = '"' + ','.join(v_cats) + '"'
            if len(v_cat_str) < 255:
                dv = DataValidation(type="list", formula1=v_cat_str, allow_blank=True)
                ws_v.add_data_validation(dv)
                dv.add(ws_v.cell(row=r, column=2))
        r += 1

        ws_v.cell(row=r, column=1, value="Filter by Subcategory:").font = blue_font
        ws_v.cell(row=r, column=1).fill = blue_bg
        if v_subcats:
            v_sub_str = '"' + ','.join(v_subcats) + '"'
            if len(v_sub_str) < 255:
                dv2 = DataValidation(type="list", formula1=v_sub_str, allow_blank=True)
                ws_v.add_data_validation(dv2)
                dv2.add(ws_v.cell(row=r, column=2))
        r += 2

        test_cols = ["Test ID", "Brand Name", "Category", "Subcategory",
                     "Total Branches", "Expected Result", "Actual Result", "Status",
                     "Notes / Defect Description"]
        header_row = r
        for c, h in enumerate(test_cols, 1):
            ws_v.cell(row=r, column=c, value=h).font = navy_font
            ws_v.cell(row=r, column=c).fill = navy
            ws_v.cell(row=r, column=c).alignment = Alignment(horizontal="center", wrap_text=True)
        r += 1

        for br in v_brands:
            ws_v.cell(row=r, column=1, value=br['Test ID']).font = reg_font
            ws_v.cell(row=r, column=1).border = border
            ws_v.cell(row=r, column=2, value=br['Brand Name']).font = reg_font
            ws_v.cell(row=r, column=2).border = border
            ws_v.cell(row=r, column=3, value=br['Category']).font = reg_font
            ws_v.cell(row=r, column=3).border = border
            ws_v.cell(row=r, column=4, value=br['Subcategory']).font = reg_font
            ws_v.cell(row=r, column=4).border = border
            ws_v.cell(row=r, column=5, value=br['Total Branches']).font = reg_font
            ws_v.cell(row=r, column=5).border = border
            ws_v.cell(row=r, column=5).alignment = Alignment(horizontal="center")
            ws_v.cell(row=r, column=6, value="Displayed").font = reg_font
            ws_v.cell(row=r, column=6).border = border
            for c in [7, 8, 9]:
                ws_v.cell(row=r, column=c).border = border
            r += 1

        last_row = r - 1

        dv_actual = DataValidation(
            type="list",
            formula1='"Displayed,Missing,Wrong Category,Wrong Subcategory,Duplicate,Wrong Country"',
            allow_blank=True)
        ws_v.add_data_validation(dv_actual)
        for row_i in range(header_row + 1, last_row + 1):
            dv_actual.add(ws_v.cell(row=row_i, column=7))

        dv_status = DataValidation(
            type="list",
            formula1='"PASS,FAIL - Missing,FAIL - Wrong Category,FAIL - Wrong Subcategory,FAIL - Count Mismatch,FAIL - Wrong Country,FAIL - Unexpected"',
            allow_blank=True)
        ws_v.add_data_validation(dv_status)
        for row_i in range(header_row + 1, last_row + 1):
            dv_status.add(ws_v.cell(row=row_i, column=8))

        ws_v.auto_filter.ref = f"A{header_row}:{get_column_letter(len(test_cols))}{last_row}"

        for col in ws_v.columns:
            max_len = max((len(str(cell.value or '')) for cell in col if len(str(cell.value or '')) < 60), default=10)
            ws_v.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 14)

    wb.save(filepath)
    return filepath


def generate_reference(file_path, verifier_names):
    print(f"\n{'='*70}")
    print(f"  GUNDRO QA REFERENCE DATA GENERATOR (v5)")
    print(f"  Source File: {os.path.basename(file_path)}")
    print(f"  Verifiers: {', '.join(verifier_names)}")
    print(f"{'='*70}")

    print(f"\n[1/3] Extracting unique brands...")
    if file_path.endswith('.xlsx'):
        brands, country_name, country_code, total_rows, total_brands, total_branches = extract_brands_from_excel(file_path)
    else:
        brands, country_name, country_code, total_rows, total_branches = extract_brands_from_sql(file_path)

    print(f"   Country: {country_name} ({country_code})")
    print(f"   Total Branch Rows: {total_rows}")
    print(f"   Unique brands extracted: {total_brands}")
    print(f"   Total branch locations: {total_branches}")

    print(f"\n[2/3] Assigning brands to verifiers (round-robin)...")
    brands, loads = assign_brands(brands, verifier_names)
    for name in verifier_names:
        v_brands = [b for b in brands if b['Assigned Verifier'] == name]
        v_branches = sum(b['Total Branches'] for b in v_brands)
        v_cats = set(b['Category'] for b in v_brands)
        print(f"   {name}: {loads[name]} brands, {v_branches} branches ({len(v_cats)} categories)")

    print(f"\n[3/3] Generating verification workbook...")
    filepath = generate_workbook(brands, verifier_names, loads, country_name, country_code,
                                  total_rows, total_brands, total_branches, file_path)

    print(f"\n{'='*70}")
    print(f"  QA VERIFICATION WORKBOOK GENERATED SUCCESSFULLY")
    print(f"{'='*70}")
    print(f"  File: {filepath}")
    print(f"{'='*70}\n")
    return filepath


def main():
    parser = argparse.ArgumentParser(description="Gundro QA Staging Reference Generator v5")
    parser.add_argument("--file", "--sql", type=str, help="Path to source Excel (.xlsx) or SQL (.sql) seed file")
    parser.add_argument("--verifiers", nargs='+', help="Verifier names")
    args = parser.parse_args()

    file_path = args.file
    verifiers = args.verifiers

    if not file_path:
        file_path = input("\nEnter path to source file (.xlsx or .sql): ").strip().strip('"')
        while not file_path or not os.path.exists(file_path):
            file_path = input("Enter valid file path: ").strip().strip('"')

    if not verifiers:
        verifiers = ["Grace", "Abdul", "Siddik"]

    generate_reference(file_path, verifiers)


if __name__ == "__main__":
    main()
